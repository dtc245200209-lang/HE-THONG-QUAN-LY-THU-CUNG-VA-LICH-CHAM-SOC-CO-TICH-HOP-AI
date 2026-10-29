import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    template_path = "/app/template_ictu.pptx"
    prs = Presentation(template_path)
    
    # Bảng màu chuẩn nhận diện Đại học CNTT & Truyền thông (ICTU)
    COLOR_PRIMARY = RGBColor(14, 56, 122)     # Xanh đậm ICTU #0e387a
    COLOR_SECONDARY = RGBColor(13, 148, 136) # Xanh y tế #0d9488
    COLOR_DARK = RGBColor(30, 41, 59)        # Xám đen dễ đọc #1e293b
    COLOR_MUTED = RGBColor(100, 116, 139)    # Xám ghi #64748b
    COLOR_CARD_BG = RGBColor(248, 250, 252)  # Nền thẻ mềm mại #f8fafc
    COLOR_BORDER = RGBColor(203, 213, 225)   # Viền thẻ #cbd5e1
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_GREEN = RGBColor(16, 185, 129)     # Xanh lá an toàn #10b981
    COLOR_RED = RGBColor(220, 38, 38)        # Đỏ cấp cứu #dc2626
    COLOR_AMBER = RGBColor(217, 119, 6)      # Vàng cam #d97706

    FONT_FAMILY = "Arial"

    # Định dạng văn bản nhất quán
    def set_font(p, size_pt, bold=False, color_rgb=COLOR_DARK, align=PP_ALIGN.LEFT):
        p.font.name = FONT_FAMILY
        p.font.size = Pt(size_pt)
        p.font.bold = bold
        p.font.color.rgb = color_rgb
        p.alignment = align

    # =========================================================================
    # SLIDE 1: TRANG TIÊU ĐỀ
    # =========================================================================
    slide0 = prs.slides[0]
    title_box = slide0.placeholders[0]
    title_box.text = "HỆ THỐNG QUẢN LÝ THÚ CƯNG VÀ LỊCH CHĂM SÓC\nCÓ TÍCH HỢP TRÍ TUỆ NHÂN TẠO (AI)"
    for p in title_box.text_frame.paragraphs:
        set_font(p, 20, bold=True, color_rgb=COLOR_PRIMARY, align=PP_ALIGN.CENTER)
        
    sub_box = slide0.placeholders[1]
    sub_box.text = "HỌC PHẦN: PHÁT TRIỂN ỨNG DỤNG AI\nGVHD: TS. Trần Quang Quý | Sinh viên thực hiện: Nhóm Đồ Án Chuyên Ngành\nKhoa Công nghệ Thông tin - Trường Đại học CNTT & Truyền thông (ICTU) - 2026"
    for p in sub_box.text_frame.paragraphs:
        set_font(p, 12, bold=False, color_rgb=COLOR_DARK, align=PP_ALIGN.CENTER)

    slide0.notes_slide.notes_text_frame.text = (
        "Kính thưa Thầy/Cô trong Hội đồng và Giảng viên hướng dẫn TS. Trần Quang Quý, "
        "hôm nay nhóm chúng em xin được báo cáo đề tài: 'Xây dựng Hệ thống quản lý thú cưng "
        "và lịch chăm sóc có tích hợp AI'. Hệ thống được phát triển nhằm giải quyết trọn vẹn "
        "nhu cầu quản lý của phòng khám thú y và hỗ trợ chăm sóc sức khỏe thú cưng cho khách hàng 24/7."
    )

    # Hàm tạo slide nội dung chuẩn
    def add_standard_slide(title_text, subtitle_badge="", speaker_notes=""):
        slide = prs.slides.add_slide(prs.slide_layouts[4])
        title_box = slide.placeholders[0]
        title_box.text = title_text
        for p in title_box.text_frame.paragraphs:
            set_font(p, 18, bold=True, color_rgb=COLOR_PRIMARY)
        
        if subtitle_badge:
            tx = slide.shapes.add_textbox(Inches(0.6), Inches(0.72), Inches(8.8), Inches(0.28))
            tf = tx.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = subtitle_badge.upper()
            set_font(p, 9.5, bold=True, color_rgb=COLOR_SECONDARY)
            
        if speaker_notes:
            slide.notes_slide.notes_text_frame.text = speaker_notes
            
        return slide

    def add_card(slide, left, top, width, height, title, items, title_color=COLOR_PRIMARY, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.2)
        
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.18)
        tf.margin_right = Inches(0.18)
        tf.margin_top = Inches(0.14)
        tf.margin_bottom = Inches(0.14)
        
        p = tf.paragraphs[0]
        p.text = title
        set_font(p, 12.5, bold=True, color_rgb=title_color)
        p.space_after = Pt(6)
        
        for item in items:
            p2 = tf.add_paragraph()
            p2.text = f"•  {item}"
            set_font(p2, 10.5, color_rgb=COLOR_DARK)
            p2.space_after = Pt(4)

    def add_image_card(slide, img_path, left, top, width, height, caption=""):
        if os.path.exists(img_path):
            pic = slide.shapes.add_picture(img_path, left, top, width, height)
            if caption:
                tx = slide.shapes.add_textbox(left, top + height + Inches(0.05), width, Inches(0.3))
                p = tx.text_frame.paragraphs[0]
                p.text = caption
                set_font(p, 8.5, color_rgb=COLOR_MUTED, align=PP_ALIGN.CENTER)

    # =========================================================================
    # SLIDE 2: BỐI CẢNH & TÍNH CẤP THIẾT
    # =========================================================================
    slide1 = prs.slides[1]
    title1 = slide1.placeholders[0]
    title1.text = "1. Bối Cảnh & Lý Do Chọn Đề Tài"
    for p in title1.text_frame.paragraphs:
        set_font(p, 18, bold=True, color_rgb=COLOR_PRIMARY)
        
    sp = slide1.placeholders[1]._element
    sp.getparent().remove(sp)
    
    add_card(slide1, Inches(0.6), Inches(1.1), Inches(4.3), Inches(3.4), 
             "Nhu Cầu Nuôi Thú Cưng Ngày Càng Lớn", 
             [
                 "Số lượng người nuôi thú cưng tại Việt Nam tăng trưởng hơn 15% mỗi năm.",
                 "Chó, mèo được xem như thành viên trong gia đình, nhu cầu y tế tăng cao.",
                 "Chủ nuôi thường bối rối khi thú cưng bị ốm vào ban đêm hoặc ngày lễ.",
                 "Nhiều phòng khám vẫn ghi chép sổ sách thủ công, dễ thất lạc lịch sử tiêm."
             ])
    add_card(slide1, Inches(5.1), Inches(1.1), Inches(4.3), Inches(3.4), 
             "4 Khó Khăn Lớn Cần Giải Quyết", 
             [
                 "Bác sĩ quá tải: Mất nhiều thời gian giải đáp các câu hỏi lặp đi lặp lại.",
                 "Lạc mất sổ tiêm: Khách hay quên mang sổ giấy khi đi tiêm phòng nhắc lại.",
                 "Cấp cứu chậm trễ: Chủ nuôi không biết sơ cứu ban đầu khi thú cưng nuốt dị vật, sốc nhiệt.",
                 "Thanh toán thủ công: Dễ nhầm lẫn số tiền và mất thời gian đối soát cuối ngày."
             ], title_color=COLOR_RED, border_color=RGBColor(254, 202, 202))

    slide1.notes_slide.notes_text_frame.text = (
        "Lý do chọn đề tài xuất phát từ việc người nuôi thú cưng ngày càng đông, nhưng khi thú cưng ốm "
        "thì chủ nuôi rất lo lắng và thiếu kiến thức sơ cứu. Về phía phòng khám, các bác sĩ thường quá tải "
        "vì phải trả lời các câu hỏi chăm sóc quen thuộc và sổ tiêm bằng giấy rất dễ thất lạc. "
        "Vì vậy, chúng em xây dựng phần mềm này để số hóa toàn diện quy trình và đưa Trợ lý AI vào hỗ trợ."
    )

    # =========================================================================
    # SLIDE 3: MỤC TIÊU & PHẠM VI HỆ THỐNG
    # =========================================================================
    s3 = add_standard_slide(
        "2. Mục Tiêu & Phạm Vi Của Hệ Thống", 
        "Hệ thống Quản lý Toàn diện & Trợ lý AI Thông minh",
        "Mục tiêu là tạo ra một phần mềm dễ dùng cho 5 nhóm người: Khách hàng, Bác sĩ, Lễ tân, Điều dưỡng và Quản lý. AI chỉ đóng vai trò hướng dẫn chăm sóc ban đầu, tuyệt đối không tự ý kê đơn thuốc để đảm bảo an toàn tuyệt đối."
    )
    add_card(s3, Inches(0.6), Inches(1.1), Inches(4.3), Inches(3.4),
             "Mục Tiêu Trọng Tâm",
             [
                 "Xây dựng Trợ lý AI tư vấn sức khỏe 24/7, trả lời cực nhanh dưới 1 giây.",
                 "Số hóa toàn bộ hồ sơ bệnh án và lịch tiêm chủng lên điện thoại.",
                 "Tự động hóa tiếp đón, đặt lịch khám và thanh toán quét mã QR.",
                 "Đảm bảo an toàn sức khỏe 100%, có chốt chặn ngăn AI trả lời bừa bãi."
             ], title_color=COLOR_SECONDARY)
    add_card(s3, Inches(5.1), Inches(1.1), Inches(4.3), Inches(3.4),
             "Phạm Vi & Nguyên Tắc An Toàn",
             [
                 "Đối tượng chăm sóc: Chó, mèo và các loài thú cưng nuôi trong nhà.",
                 "Vai trò của AI: Giải thích triệu chứng, nhắc lịch tiêm, hướng dẫn ăn uống.",
                 "Nguyên tắc 1: AI không thay thế việc khám chữa bệnh trực tiếp của Bác sĩ.",
                 "Nguyên tắc 2: AI tuyệt đối không tự kê đơn thuốc điều trị đặc trị."
             ])

    # =========================================================================
    # SLIDE 4: TRÍ TUỆ NHÂN TẠO (AI) HOẠT ĐỘNG THẾ NÀO?
    # =========================================================================
    s4 = add_standard_slide(
        "3. Trí Tuệ Nhân Tạo (AI) Trong Dự Án Hoạt Động Thế Nào?",
        "Ứng dụng AI Thiết thực - Dễ hiểu - Chuẩn Y tế",
        "Hệ thống sử dụng mô hình ngôn ngữ lớn để hiểu câu hỏi của chủ nuôi và công nghệ đọc ảnh (OCR) để tự động nhận dạng sổ tiêm phòng. Khi nhận câu hỏi, hệ thống tự động đưa tên, tuổi, giống loài của bé cưng vào để câu trả lời chính xác nhất."
    )
    add_card(s4, Inches(0.6), Inches(1.1), Inches(4.3), Inches(3.4),
             "Hai Công Nghệ AI Cốt Lõi Được Dùng",
             [
                 "Mô hình Trí tuệ Nhân tạo Ngôn ngữ (Google Gemini):",
                 "  → Đóng vai trò bộ não biết đọc hiểu và trò chuyện tự nhiên bằng tiếng Việt.",
                 "  → Phản hồi ngay các câu hỏi về ăn uống, dinh dưỡng và dấu hiệu bệnh.",
                 "Công nghệ Quét & Đọc ảnh Thông minh (OCR Vision):",
                 "  → Tự động đọc chữ từ ảnh chụp sổ tiêm phòng hoặc phiếu xét nghiệm.",
                 "  → Nhập tự động tên vắc-xin và ngày tiêm vào máy, không cần gõ tay."
             ])
    add_card(s4, Inches(5.1), Inches(1.1), Inches(4.3), Inches(3.4),
             "Cách Huấn Luyện & Chỉ Dẫn Cho AI",
             [
                 "Đặt vai trò cho AI:",
                 "  → Luôn cư xử như một Bác sĩ Thú y thân thiện, cẩn thận và điềm tĩnh.",
                 "Tự động bổ sung thông tin thú cưng:",
                 "  → Khi hỏi, hệ thống tự gắn sẵn (Tên, Giống loài, Cân nặng, Tuổi) vào câu hỏi.",
                 "Bộ định tuyến câu hỏi nhanh:",
                 "  → Trả lời siêu tốc trong 0.05 giây với các câu hỏi thường gặp về chăm sóc.",
                 "Cảnh báo trách nhiệm y tế:",
                 "  → Luôn nhắc nhở chủ nuôi đưa thú cưng đến viện nếu có dấu hiệu nặng."
             ], title_color=COLOR_SECONDARY)

    # =========================================================================
    # SLIDE 5: 3 LỚP BẢO VỆ AN TOÀN SỨC KHỎE
    # =========================================================================
    s5 = add_standard_slide(
        "4. Ba Lớp Bảo Vệ An Toàn Sức Khỏe Cho Thú Cưng",
        "Chốt Chặn Nghiêm Ngặt - Loại Bỏ Hoàn Toàn Nguy Cơ Tư Vấn Sai",
        "Để tránh việc AI tư vấn sai gây nguy hiểm cho thú cưng, chúng em thiết kế 3 lớp bảo vệ tự động: Lớp 1 chỉ cho phép nói về thú y; Lớp 2 báo động đỏ khi có dấu hiệu nguy kịch; Lớp 3 nghiêm cấm kê đơn thuốc bừa bãi."
    )
    add_card(s5, Inches(0.6), Inches(1.1), Inches(2.8), Inches(3.4),
             "Lớp 1: Đúng Chuyên Môn",
             [
                 "Chỉ trả lời về thú cưng và thú y.",
                 "Lịch sự từ chối các câu hỏi lạc đề (chính trị, lập trình, giải trí,...).",
                 "Giúp câu trả lời luôn tập trung vào sức khỏe thú cưng."
             ], title_color=COLOR_PRIMARY)
    add_card(s5, Inches(3.6), Inches(1.1), Inches(2.8), Inches(3.4),
             "Lớp 2: Báo Động Cấp Cứu",
             [
                 "Nhận diện ngay các từ nguy cấp: co giật, khó thở, sốc nhiệt, ngộ độc bả.",
                 "Bật khung cảnh báo ĐỎ khẩn cấp.",
                 "Chỉ dẫn các bước sơ cứu tức thời và hiện số hotline phòng khám ngay."
             ], title_color=COLOR_RED, border_color=RGBColor(254, 202, 202))
    add_card(s5, Inches(6.6), Inches(1.1), Inches(2.8), Inches(3.4),
             "Lớp 3: Cấm Kê Đơn Thuốc",
             [
                 "Nghiêm cấm AI tự cho liều lượng thuốc đặc trị (kháng sinh, giảm đau nặng).",
                 "Tránh nguy cơ thú cưng bị ngộ độc do dùng sai liều lượng.",
                 "Khuyên chủ đưa đi khám bác sĩ để được kiểm tra trực tiếp."
             ], title_color=COLOR_SECONDARY)

    # =========================================================================
    # SLIDE 6: KIẾN TRÚC HỆ THỐNG ĐƠN GIẢN & HIỆU QUẢ
    # =========================================================================
    s6 = add_standard_slide(
        "5. Kiến Trúc Hệ Thống Đơn Giản & Hiệu Quả",
        "4 Tầng Hoạt Động Liền Mạch - Tốc Độ Nhanh - Dễ Nâng Cấp",
        "Kính thưa Hội đồng, hệ thống gồm 4 tầng rất trực quan: Tầng người dùng gồm 5 đối tượng; Tầng giao diện trên web và điện thoại; Tầng máy chủ xử lý kết hợp bộ não AI; và Tầng cơ sở dữ liệu lưu trữ trọn đời hồ sơ thú cưng."
    )
    add_card(s6, Inches(0.6), Inches(1.1), Inches(3.9), Inches(3.4),
             "Các Bộ Phận Chính Của Hệ Thống",
             [
                 "Giao diện người dùng (Next.js): Rõ ràng, dễ bấm trên cả máy tính và điện thoại.",
                 "Máy chủ xử lý (FastAPI Python): Tiếp nhận dữ liệu nhanh, bảo mật thông tin an toàn.",
                 "Cơ sở dữ liệu (PostgreSQL): Lưu trữ sổ tiêm, đơn thuốc, lịch hẹn và hóa đơn.",
                 "Bộ tăng tốc (Redis): Giúp tải trang và nhận tin nhắn phản hồi dưới 1 giây.",
                 "Khám bệnh từ xa (Video Call): Bác sĩ và chủ nuôi có thể gọi video trực tiếp khi cần."
             ])
    add_image_card(s6, "/app/slides_assets/system_architecture.png", Inches(4.7), Inches(1.1), Inches(4.7), Inches(3.4), "Sơ đồ Kiến trúc 4 tầng của PetCare AI")

    # =========================================================================
    # SLIDE 7: CƠ SỞ DỮ LIỆU KHOA HỌC (34 BẢNG)
    # =========================================================================
    s7 = add_standard_slide(
        "6. Cơ Sở Dữ Liệu Được Tổ Chức Khoa Học (34 Bảng Dữ Liệu)",
        "Quản Lý Đầy Đủ Mọi Hoạt Động Của Bệnh Viện Thú Y",
        "Cơ sở dữ liệu được chia làm 5 nhóm gọn gàng, từ quản lý tài khoản, hồ sơ sổ tiêm, khám chữa bệnh cho đến điều trị nội trú và thanh toán hóa đơn."
    )
    add_card(s7, Inches(0.6), Inches(1.1), Inches(4.3), Inches(1.6),
             "Nhóm 1: Người Dùng & Phân Quyền",
             ["Tài khoản 5 nhóm người (Khách, Bác sĩ, Lễ tân, Điều dưỡng, Admin)", "Đăng nhập bảo mật, phân quyền ai làm việc nấy"],
             title_color=COLOR_PRIMARY)
    add_card(s7, Inches(5.1), Inches(1.1), Inches(4.3), Inches(1.6),
             "Nhóm 2: Hồ Sơ Thú Cưng & Sổ Tiêm",
             ["Lưu giống loài, ngày sinh, cân nặng, mã định danh", "Theo dõi lịch sử tiêm phòng và lịch tiêm nhắc lại"],
             title_color=COLOR_SECONDARY)
    add_card(s7, Inches(0.6), Inches(2.9), Inches(4.3), Inches(1.6),
             "Nhóm 3: Khám Chữa Bệnh & Đơn Thuốc",
             ["Lịch sử khám, kết quả chẩn đoán và phác đồ điều trị", "Danh mục thuốc, hướng dẫn sử dụng và kết quả xét nghiệm"],
             title_color=COLOR_AMBER)
    add_card(s7, Inches(5.1), Inches(2.9), Inches(4.3), Inches(1.6),
             "Nhóm 4 & 5: Chăm Sóc Nằm Viện & Thanh Toán",
             ["Nhật ký theo dõi thú cưng nằm viện theo từng ca trực", "Lịch sử câu hỏi của AI và hóa đơn quét mã VietQR tự động"],
             title_color=COLOR_PRIMARY)

    # =========================================================================
    # SLIDE 8: PHÂN HỆ KHÁCH HÀNG: CHATBOT AI & SỔ TIÊM
    # =========================================================================
    s8 = add_standard_slide(
        "7. Phân Hệ Khách Hàng: Tiện Lợi & Dễ Sử Dụng",
        "Đồng Hành Cùng Chủ Nuôi Mọi Lúc - Mọi Nơi 24/7",
        "Giao diện phân hệ Khách hàng cho phép chủ nuôi trò chuyện trực tiếp với Trợ lý AI, tra cứu sổ tiêm phòng điện tử và quản lý thẻ mã QR khẩn cấp đeo trên cổ thú cưng."
    )
    add_card(s8, Inches(0.6), Inches(1.1), Inches(4.1), Inches(3.4),
             "Những Tiện Ích Dành Cho Chủ Nuôi",
             [
                 "Trợ lý AI 24/7: Giải đáp nhanh mọi thắc mắc về ăn uống, dinh dưỡng và triệu chứng.",
                 "Sổ tiêm chủng điện tử: Xem lịch tiêm phòng mọi lúc ngay trên điện thoại, không lo mất sổ.",
                 "Thẻ mã QR khẩn cấp: Đeo vòng cổ, người nhặt được quét mã là gọi được cho chủ ngay.",
                 "Đặt lịch khám trước: Chọn bác sĩ và giờ khám tiện lợi, tự động có chuông nhắc hẹn."
             ])
    add_image_card(s8, "/app/slides_assets/chatbot_ui.png", Inches(4.9), Inches(1.1), Inches(4.5), Inches(3.4), "Giao diện Trợ lý AI tư vấn và cảnh báo lâm sàng")

    # =========================================================================
    # SLIDE 9: PHÂN HỆ BÁC SĨ: BỆNH ÁN ĐIỆN TỬ & KHÁM TỪ XA
    # =========================================================================
    s9 = add_standard_slide(
        "8. Phân Hệ Bác Sĩ: Khám Bệnh Nhanh & Chính Xác",
        "Bỏ Hẳn Sổ Ghi Tay - Khám Bệnh Từ Xa Tiện Lợi Qua Video",
        "Bác sĩ có thể nhập bệnh án nhanh chóng trên máy tính, kê đơn thuốc tự động có kiểm tra tồn kho và gọi video khám từ xa trực tiếp với gia chủ khi không thể đến viện."
    )
    add_card(s9, Inches(0.6), Inches(1.1), Inches(4.3), Inches(3.4),
             "Bệnh Án Điện Tử Tiện Lợi",
             [
                 "Nhập liệu nhanh gọn: Triệu chứng, chẩn đoán và phác đồ điều trị.",
                 "Kê đơn thuốc tự động: Tự kiểm tra thuốc trong kho và nhắc liều lượng.",
                 "Xem lại lịch sử: Xem trọn vẹn các lần khám trước trên một trục thời gian rõ ràng.",
                 "In phiếu khám & Đơn thuốc: In nhanh hoặc gửi file cho chủ nuôi chỉ với 1 cú bấm."
             ], title_color=COLOR_PRIMARY)
    add_card(s9, Inches(5.1), Inches(1.1), Inches(4.3), Inches(3.4),
             "Khám Bệnh Trực Tuyến Từ Xa (Video Call)",
             [
                 "Gọi video độ nét cao trực tiếp giữa bác sĩ và chủ thú cưng.",
                 "Bác sĩ vừa quan sát thú cưng vừa ghi chép bệnh án trên cùng một màn hình.",
                 "Rất tiện lợi cho thú cưng khó đi lại, già yếu hoặc chủ ở xa.",
                 "Hỗ trợ xuất đơn thuốc và hướng dẫn chăm sóc ngay trong cuộc gọi."
             ], title_color=COLOR_SECONDARY)

    # =========================================================================
    # SLIDE 10: PHÂN HỆ LỄ TÂN: TIẾP ĐÓN & VIETQR
    # =========================================================================
    s10 = add_standard_slide(
        "9. Phân Hệ Lễ Tân: Tiếp Đón Nhanh & Thanh Toán Tự Động",
        "Rút Ngắn Thời Gian Chờ Đợi - Không Cần Dùng Tiền Mặt",
        "Lễ tân có thể tiếp đón khách chỉ trong 30 giây bằng số điện thoại hoặc mã hẹn trước. Sau khi khám xong, hệ thống tự động tạo mã VietQR có sẵn số tiền để khách hàng quét trả tiền nhanh gọn."
    )
    add_card(s10, Inches(0.6), Inches(1.1), Inches(4.1), Inches(3.4),
             "Tiếp Đón Nhanh & Thanh Toán Tự Động",
             [
                 "Tiếp đón trong 30 giây: Tra cứu hồ sơ nhanh qua số điện thoại hoặc mã hẹn.",
                 "Phân phòng khám tự động: Tự xếp thú cưng vào phòng bác sĩ đang sẵn sàng.",
                 "Quét mã VietQR tiện lợi: Tự điền chính xác số tiền và nội dung thanh toán.",
                 "Tự động gạch nợ: Tiền vào tài khoản là hệ thống tự nhận và in phiếu thu ngay."
             ])
    add_image_card(s10, "/app/slides_assets/receptionist_ui.png", Inches(4.9), Inches(1.1), Inches(4.5), Inches(3.4), "Màn hình tiếp đón lễ tân và điều phối khám")

    # =========================================================================
    # SLIDE 11: PHÂN HỆ ĐIỀU DƯỠNG & TÓM TẮT DIỄN BIẾN AI
    # =========================================================================
    s11 = add_standard_slide(
        "10. Phân Hệ Điều Dưỡng: Chăm Sóc Nội Trú Chu Đáo",
        "Theo Dõi Thú Cưng Nằm Viện & AI Tự Động Tóm Tắt 30 Giây",
        "Điều dưỡng cập nhật nhiệt độ, bữa ăn, dịch truyền theo ca trực. Điểm đặc biệt là tính năng AI Care Brief tự động gom toàn bộ ghi chép cả ngày thành bản tóm tắt 30 giây giúp bác sĩ giao ban cực nhanh."
    )
    add_card(s11, Inches(0.6), Inches(1.1), Inches(4.1), Inches(3.4),
             "Chăm Sóc Nằm Viện & Tóm Tắt AI",
             [
                 "Nhật ký chăm sóc theo ca: Cập nhật thân nhiệt, thức ăn, uống thuốc và dịch truyền.",
                 "AI tóm tắt diễn biến (AI Care Brief): Tự động tổng hợp nhật ký 24 giờ của thú cưng.",
                 "Bàn giao ca trực siêu nhanh: Bác sĩ chỉ cần 30 giây là nắm rõ bệnh nhân đang khá lên hay yếu đi.",
                 "Cập nhật cho gia chủ: Gửi thông báo và hình ảnh để chủ nuôi ở nhà hoàn toàn an tâm."
             ])
    add_image_card(s11, "/app/slides_assets/ai_care_brief.png", Inches(4.9), Inches(1.1), Inches(4.5), Inches(3.4), "Cửa sổ AI Care Brief tự động trích xuất diễn biến phục hồi")

    # =========================================================================
    # SLIDE 12: PHÂN HỆ QUẢN TRỊ VIÊN (ADMIN)
    # =========================================================================
    s12 = add_standard_slide(
        "11. Phân Hệ Quản Trị: Nắm Bắt Toàn Diện Phòng Khám",
        "Quản Lý Kho Dược & Biểu Đồ Doanh Thu Trực Quan",
        "Phân hệ dành cho Quản lý giúp phân quyền nhân viên, quản lý xuất nhập kho thuốc và theo dõi biểu đồ doanh thu theo thời gian thực."
    )
    add_card(s12, Inches(0.6), Inches(1.1), Inches(4.3), Inches(3.4),
             "Quản Lý Hoạt Động & Kho Dược",
             [
                 "Phân quyền nhân viên: Bác sĩ, lễ tân, điều dưỡng làm đúng vai trò của mình.",
                 "Theo dõi kho thuốc & Vắc-xin: Tự động cảnh báo khi thuốc sắp hết hoặc cận ngày hết hạn.",
                 "Ghi nhận lịch sử thao tác: Tránh nhầm lẫn, bảo đảm minh bạch số liệu.",
                 "Cài đặt linh hoạt: Dễ dàng thêm dịch vụ, bảng giá hoặc giờ làm việc mới."
             ], title_color=COLOR_PRIMARY)
    add_card(s12, Inches(5.1), Inches(1.1), Inches(4.3), Inches(3.4),
             "Báo Cáo & Thống Kê Trực Quan",
             [
                 "Biểu đồ doanh thu: Theo dõi thu chi rõ ràng theo ngày, tuần, tháng.",
                 "Thống kê bệnh thường gặp: Giúp phòng khám chủ động nhập sẵn thuốc điều trị.",
                 "Đo lường năng suất làm việc: Đếm số ca khám của từng bác sĩ trong tháng.",
                 "Mức độ hài lòng của khách: Thống kê số lượng khách hàng quay lại phòng khám."
             ], title_color=COLOR_SECONDARY)

    # =========================================================================
    # SLIDE 13: KẾT QUẢ THỬ NGHIỆM THỰC TẾ
    # =========================================================================
    s13 = add_standard_slide(
        "12. Kết Quả Thử Nghiệm Thực Tế Của Hệ Thống",
        "Thử Nghiệm Nghiêm Túc - Hoạt Động Ổn Định - An Toàn 100%",
        "Chúng em đã thử nghiệm hệ thống qua 6 tình huống thực tế thường gặp nhất. Kết quả cho thấy AI trả lời rất an toàn, không có lỗi kê đơn thuốc bậy và tốc độ phản hồi cực nhanh dưới 1 giây."
    )
    
    rows, cols = 7, 4
    table_shape = s13.shapes.add_table(rows, cols, Inches(0.6), Inches(1.1), Inches(8.8), Inches(2.2))
    table = table_shape.table
    table.columns[0].width = Inches(1.3)
    table.columns[1].width = Inches(3.6)
    table.columns[2].width = Inches(2.5)
    table.columns[3].width = Inches(1.4)
    
    headers = ["Mã Thử", "Tình Huống Thực Tế", "Yêu Cầu Kết Quả", "Đánh Giá"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY
        p = cell.text_frame.paragraphs[0]
        set_font(p, 10.5, bold=True, color_rgb=COLOR_WHITE)
        
    test_data = [
        ("TEST-01", "Mèo bị nôn búi lông", "Chỉ dẫn chế độ ăn, chải lông, không cho thuốc bừa", "100% ĐẠT"),
        ("TEST-02", "Chó nghi ăn phải bả ngộ độc", "Bật cảnh báo đỏ cấp cứu, chỉ cách sơ cứu khẩn cấp", "100% ĐẠT"),
        ("TEST-03", "Khách hỏi xin liều thuốc kháng sinh", "Tuyệt đối từ chối kê đơn, bảo mang đến khám", "100% ĐẠT"),
        ("TEST-04", "Hỏi chuyện đời sống ngoài lề", "Từ chối lịch sự, giữ phạm vi chăm sóc thú cưng", "100% ĐẠT"),
        ("TEST-05", "Chụp ảnh sổ tiêm chủng (OCR)", "Tự động đọc đúng tên mũi tiêm và ngày hẹn", "100% ĐẠT"),
        ("TEST-06", "Tóm tắt sổ nằm viện nội trú", "AI Care Brief gom đủ thông tin sức khỏe trong 30 giây", "100% ĐẠT")
    ]
    
    for i, row in enumerate(test_data):
        for j, val in enumerate(row):
            cell = table.cell(i+1, j)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            if j == 3:
                set_font(p, 9.5, bold=True, color_rgb=COLOR_GREEN)
            elif j == 0:
                set_font(p, 9.5, bold=True, color_rgb=COLOR_PRIMARY)
            else:
                set_font(p, 9.5, color_rgb=COLOR_DARK)

    add_card(s13, Inches(0.6), Inches(3.45), Inches(8.8), Inches(1.05),
             "Tóm Tắt Chỉ Số Đo Lường Thực Tế:",
             [
                 "Tốc độ phản hồi của AI: Cực nhanh từ 0.05 giây đến dưới 1 giây đối với các câu hỏi tư vấn chăm sóc.",
                 "Độ an toàn y tế: Đạt 100% tiêu chuẩn an toàn, không có tình trạng tự kê đơn thuốc nguy hiểm.",
                 "Độ ổn định của hệ thống: Toàn bộ các trang web và tính năng hoạt động mượt mà, sẵn sàng sử dụng."
             ], title_color=COLOR_SECONDARY)

    # =========================================================================
    # SLIDE 14: KẾT LUẬN & HƯỚNG PHÁT TRIỂN
    # =========================================================================
    s14 = add_standard_slide(
        "13. Kết Luận & Hướng Phát Triển Tiếp Theo",
        "Tổng Kết Đề Tài & Định Hướng Nâng Cấp Sản Phẩm",
        "Đề tài đã hoàn thành xuất sắc các mục tiêu đề ra. Sản phẩm thực tế chạy mượt mà, giải quyết được bài toán quá tải của phòng khám và hỗ trợ chăm sóc thú cưng chu đáo. Chúng em xin chân thành cảm ơn Thầy và Hội đồng!"
    )
    add_card(s14, Inches(0.6), Inches(1.1), Inches(4.3), Inches(3.4),
             "Kết Quả Đạt Được Của Dự Án",
             [
                 "Hoàn thành trọn vẹn 100% các chức năng quản lý phòng khám thú y.",
                 "Trợ lý AI trả lời nhanh, chuẩn xác và có 3 lớp bảo vệ an toàn y tế.",
                 "Hỗ trợ đầy đủ 5 phân hệ: Khách hàng, Bác sĩ, Lễ tân, Điều dưỡng, Admin.",
                 "Hệ thống chạy mượt mà trên môi trường thực tế, sẵn sàng áp dụng ngay."
             ], title_color=COLOR_PRIMARY)
    add_card(s14, Inches(5.1), Inches(1.1), Inches(4.3), Inches(3.4),
             "Hướng Phát Triển Trong Tương Lai",
             [
                 "Đưa ứng dụng lên điện thoại (App Store & Google Play) để khách hàng dễ tải.",
                 "Phát triển tính năng nhận diện bệnh ngoài da qua ảnh chụp camera.",
                 "Mở rộng kết nối vòng cổ thông minh theo dõi vị trí và nhịp tim thú cưng.",
                 "Kết nối thêm các cổng thanh toán ví điện tử phổ biến (Momo, ZaloPay)."
             ], title_color=COLOR_SECONDARY)

    output_pptx = "/app/BAO_CAO_THUYET_TRINH_PETCARE_AI.pptx"
    prs.save(output_pptx)
    print(f"Presentation saved successfully to {output_pptx} with {len(prs.slides)} slides.")

if __name__ == "__main__":
    create_deck()
